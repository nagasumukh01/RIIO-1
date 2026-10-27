#!/usr/bin/env python3
"""Create the assignment Word report from saved real experiment artifacts."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            if row_index == 0:
                set_cell_shading(cell, "1F4E79")
                for run in cell.paragraphs[0].runs:
                    run.font.color.rgb = None
                    run.bold = True
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(document, text: str, level: int = 1):
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_bullet(document, text: str):
    document.add_paragraph(text, style="List Bullet")


def metric(value, digits: int = 4) -> str:
    if value is None:
        return "Pending"
    return f"{value:.{digits}f}"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", default=str(ROOT / "assignments" / "Week_8_Project_Report_RobustDeep.docx"))
    args = parser.parse_args()

    model_dir = ROOT / "artifacts" / "models" / "ff_c23_subset10"
    summary = json.loads((model_dir / "training_summary.json").read_text(encoding="utf-8"))
    evaluation = json.loads((model_dir / "evaluations" / "faceforensics_test" / "evaluation.json").read_text(encoding="utf-8"))
    test = evaluation["metrics"]
    val = summary["val_metrics"]
    vit_dir = ROOT / "artifacts" / "models" / "ff_c23_vit_small_milestone"
    vit_summary = json.loads((vit_dir / "training_summary.json").read_text(encoding="utf-8"))
    vit_val = vit_summary["val_metrics"]
    vit_test = vit_summary["test_metrics"]
    figures = model_dir / "evaluations" / "faceforensics_test"
    explainability = ROOT / "artifacts" / "figures" / "explainability" / "ff_c23_subset10"
    explanation_figure = next(iter(sorted(explainability.glob("*.png"))), None)

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("WEEK 8 PROJECT REPORT")
    run.bold = True
    run.font.size = Pt(18)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("RobustDeep: Real-Data FaceForensics++ Training and Evaluation")
    run.bold = True
    run.font.size = Pt(14)
    for text in [
        "Vertical: A - Group I",
        "Prepared by: Anvitha GR [R23EF028]",
        "Bhoomika S [R23EF048]",
        "D Nagasumukh [R23EF059]",
        "Guide: Prof. Amarjit Khan | Center for Artificial Intelligence & Machine Learning",
        "Project repository: RobustDeep local workspace",
    ]:
        paragraph = document.add_paragraph(text)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph()

    add_heading(document, "Executive Summary")
    document.add_paragraph(
        "RobustDeep is a real-data deepfake detection pipeline for FaceForensics++, Celeb-DF v2, and DFDC. "
        "This report records the first verified real-data execution: an approved FaceForensics++ c23 subset was downloaded, "
        "preprocessed with video-level splits, and evaluated with EfficientNet-B0. The best saved checkpoint produced "
        f"validation AUROC {metric(val['auroc'])} and held-out test AUROC {metric(test['auroc'])}. "
        "The test set contains only 8 videos, so these values are preliminary and must not be generalized to the full benchmark."
    )
    document.add_paragraph(
        "No Celeb-DF or DFDC videos were present in this run. Their metrics are therefore pending. Synthetic data remains "
        "restricted to the Stage-1 smoke test and is not included in the results below."
    )

    add_heading(document, "Assignment Requirements and Completion Status")
    table = document.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ["Requirement", "Evidence", "Status"]):
        set_cell_text(cell, text, True)
    rows = [
        ("Prepare the three official datasets", "Folder structure and official acquisition instructions are implemented.", "Ready; data-dependent"),
        ("Preprocess FaceForensics++", "50 approved c23 videos imported; 35 train, 7 validation, 8 test videos yielded face crops.", "Completed for subset"),
        ("Train a real CNN baseline", "EfficientNet-B0 checkpoint and real evaluation artifacts saved under artifacts/models.", "Completed; interrupted 5-epoch run"),
        ("Train an initial ViT/hybrid model", "ViT-Small trained for one epoch on the same real FF++ split with saved evaluation artifacts.", "Completed for subset"),
        ("Train Celeb-DF and DFDC", "No local videos or face crops were available for these datasets.", "Pending data"),
        ("Synthetic isolation", "train.py rejects synthetic; smoke_test.py remains the synthetic entry point.", "Verified"),
    ]
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            set_cell_text(cell, value)
    style_table(table)

    add_heading(document, "Preprocessing Pipeline")
    document.add_paragraph(
        "The official FaceForensics++ downloader was used with c23 compression and a 10-video request per source group. "
        "The downloaded original, Deepfakes, Face2Face, FaceSwap, and NeuralTextures videos were copied into the project’s "
        "expected raw folders. OpenCV Haar detection was selected because RetinaFace and MTCNN were not available in the active environment."
    )
    document.add_paragraph("Observed preprocessing output:")
    add_bullet(document, "Dataset: FaceForensics++ c23 subset")
    add_bullet(document, "Raw videos imported: 50 across five project-supported groups")
    add_bullet(document, "Face crops: 35 train videos, 7 validation videos, 8 test videos")
    add_bullet(document, "Frames per video: 16; face crop size: 224 x 224 pixels")
    add_bullet(document, "Manifests: data/processed/faceforensics/{train,val,test}_manifest.json")

    add_heading(document, "Model and Training Setup")
    document.add_paragraph(
        "The evaluated model is EfficientNet-B0 with pretrained ImageNet weights, mean temporal aggregation, AdamW optimization, "
        "random seed 42, and real-data-only loading. The checkpoint is stored separately from synthetic smoke-test checkpoints at:"
    )
    document.add_paragraph(str(model_dir / "best.pt"))
    document.add_paragraph("Initial ViT-Small checkpoint:")
    document.add_paragraph(str(vit_dir / "best.pt"))
    document.add_paragraph(
        "The five-epoch CPU run completed through epoch 3 before termination. The best checkpoint was selected using validation AUROC; "
        "the final test set was evaluated once without target threshold tuning. The saved report marks the run as incomplete."
    )

    add_heading(document, "Saved Preliminary Evaluation Results")
    document.add_paragraph("All values below are read from the saved real-data JSON artifacts and are labelled REAL DATA.")
    table = document.add_table(rows=1, cols=7)
    headers = ["Dataset", "Model", "Split", "Samples", "Accuracy", "AUROC", "F1 / EER"]
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_text(cell, text, True)
    values = [
        ["FaceForensics++ c23", "EfficientNet-B0", "Validation", "7", metric(val["accuracy"]), metric(val["auroc"]), f"{metric(val['f1'])} / {metric(val['eer'])}"],
        ["FaceForensics++ c23", "EfficientNet-B0", "Test", str(test["n_videos"]), metric(test["accuracy"]), metric(test["auroc"]), f"{metric(test['f1'])} / {metric(test['eer'])}"],
        ["FaceForensics++ c23", "ViT-Small", "Validation", str(vit_val["n_videos"]), metric(vit_val["accuracy"]), metric(vit_val["auroc"]), f"{metric(vit_val['f1'])} / {metric(vit_val['eer'])}"],
        ["FaceForensics++ c23", "ViT-Small", "Test", str(vit_test["n_videos"]), metric(vit_test["accuracy"]), metric(vit_test["auroc"]), f"{metric(vit_test['f1'])} / {metric(vit_test['eer'])}"],
        ["Celeb-DF v2", "EfficientNet-B0", "Test", "Pending", "Pending", "Pending", "Pending"],
        ["DFDC", "EfficientNet-B0", "Test", "Pending", "Pending", "Pending", "Pending"],
    ]
    for row_values in values:
        cells = table.add_row().cells
        for cell, value in zip(cells, row_values):
            set_cell_text(cell, value)
    style_table(table)
    document.add_paragraph(
        "The FF++ test confusion matrix is TN=0, FP=2, FN=1, TP=5 at the fixed 0.5 decision threshold. "
        "Because the test sample is small and imbalanced, the metrics are evidence of this subset only, not a full-dataset benchmark claim."
    )

    add_heading(document, "Training Log Evidence")
    document.add_paragraph(
        "The best validation checkpoint was saved after the second completed epoch with validation AUROC 0.7000 and validation accuracy 0.8571. "
        "A later completed epoch had lower validation AUROC, so the earlier checkpoint was retained. The final export was interrupted, "
        "but the saved checkpoint was independently evaluated and its test artifacts were preserved."
    )
    document.add_paragraph(
        f"The initial ViT-Small run completed one epoch with validation AUROC {metric(vit_val['auroc'])} and test AUROC {metric(vit_test['auroc'])}. "
        "The validation/test gap is recorded as a failure-analysis result, not hidden by reporting validation performance as the final benchmark."
    )

    add_heading(document, "Results Figures")
    roc = figures / "roc_test.png"
    confusion = figures / "confusion_matrix_test.png"
    vit_roc = vit_dir / "roc_test.png"
    if roc.exists():
        document.add_picture(str(roc), width=Inches(5.5))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.add_paragraph("Figure 1. REAL DATA FaceForensics++ test ROC curve.").style = "Caption"
    if confusion.exists():
        document.add_picture(str(confusion), width=Inches(4.5))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.add_paragraph("Figure 2. REAL DATA FaceForensics++ test confusion matrix at threshold 0.5.").style = "Caption"
    if vit_roc.exists():
        document.add_picture(str(vit_roc), width=Inches(5.5))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.add_paragraph("Figure 3. REAL DATA FaceForensics++ ViT-Small test ROC curve.").style = "Caption"
    if explanation_figure is not None:
        document.add_picture(str(explanation_figure), width=Inches(5.5))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.add_paragraph("Figure 4. Model attribution visualization; attribution is not forensic proof.").style = "Caption"

    add_heading(document, "Limitations and Next Steps")
    for item in [
        "Complete a stable multi-epoch run on a GPU or a smaller controlled subset, then compare only using the untouched test split.",
        "The ViT-Small validation/test gap is a failure-analysis result; do not report the validation AUROC as its final test score.",
        "Acquire and preprocess licensed Celeb-DF v2 and DFDC subsets before reporting cross-dataset or three-dataset results.",
        "Increase the number and identity diversity of videos; the current FF++ test set has only 8 videos.",
        "Run the six-pair matrix only after source and target face crops exist for all three official datasets.",
        "Do not treat the validation AUROC as the final benchmark result; the held-out test AUROC is the relevant reported value.",
        "Keep synthetic smoke-test metrics in the explicitly labelled non-research dashboard section.",
    ]:
        add_bullet(document, item)

    add_heading(document, "Reproducibility and Repository")
    document.add_paragraph("Primary commands used for this report:")
    for command in [
        "python scripts/preprocess_faces.py --dataset faceforensics --max-videos 50 --frames-per-video 16",
        "python scripts/train.py --dataset faceforensics --model efficientnet_b0 --epochs 5 --max-videos 50 --experiment-id ff_c23_subset10",
        "python scripts/train.py --dataset faceforensics --model vit_small --epochs 1 --max-videos 50 --num-workers 0 --experiment-id ff_c23_vit_small_milestone",
        "python scripts/evaluate.py --dataset faceforensics --split test --model efficientnet_b0 --checkpoint artifacts/models/ff_c23_subset10/best.pt",
        "python scripts/run_explainability.py --checkpoint artifacts/models/ff_c23_subset10/best.pt --dataset faceforensics --split test --model efficientnet_b0 --max-videos 8",
    ]:
        paragraph = document.add_paragraph(style="No Spacing")
        run = paragraph.add_run(command)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
    document.add_paragraph("The dashboard reads saved real outputs from artifacts/models and is available at http://localhost:8501 when Streamlit is running.")

    add_heading(document, "References")
    references = [
        "FaceForensics++ official repository and access instructions: https://github.com/ondyari/FaceForensics",
        "FaceForensics++ access form: https://docs.google.com/forms/d/e/1FAIpQLSdRRR3L5zAv6tQ_CKxmK4W96tAab_pfBu2EKAgQbeDVhmXagg/viewform",
        "Rössler et al., FaceForensics++: Learning to Detect Manipulated Facial Images, ICCV 2019: https://arxiv.org/abs/1901.08971",
        "RobustDeep dataset protocol: docs/dataset_protocol.md",
        "RobustDeep experiment protocol: docs/experiment_protocol.md",
    ]
    for reference in references:
        add_bullet(document, reference)

    add_heading(document, "Conclusion")
    document.add_paragraph(
        "The project now demonstrates a reproducible real-data FaceForensics++ path from approved download through face preprocessing, "
        "checkpoint storage, held-out evaluation, visual artifacts, and explainability outputs. The saved results are useful preliminary evidence, "
        "but they are limited by the small subset and interrupted training run. Celeb-DF v2, DFDC, and the complete six-pair research matrix remain pending until their licensed data are supplied and processed."
    )

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(output)


if __name__ == "__main__":
    main()
